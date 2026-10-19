"""
core/lp_delete.py: LP_Delete core engine with XML zip surgery, openpyxl fallback, and verification.
"""

import io
import math
import re
import time
import xml.etree.ElementTree as ET
import zipfile
from typing import Any, Callable, Dict, List, Optional, Tuple

import openpyxl
import pandas as pd

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_ROW_RE = re.compile(r'<row r="(\d+)"([^>]*?)(/>|>(.*?)</row>)', re.DOTALL)


def _col_letters(n: int) -> str:
    letters = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _escape_xml_text(s: str) -> str:
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _find_sheet_xml_path(zf: zipfile.ZipFile, sheet_name: str) -> str:
    root = ET.fromstring(zf.read("xl/workbook.xml"))
    rid = None
    sheets_el = root.find(f"{{{NS_MAIN}}}sheets")
    if sheets_el is not None:
        for sheet_el in sheets_el.findall(f"{{{NS_MAIN}}}sheet"):
            if sheet_el.get("name") == sheet_name:
                rid = sheet_el.get(f"{{{NS_R}}}id")
                break
    if rid is None:
        raise ValueError(f"Sheet '{sheet_name}' not found in workbook.xml")

    rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    target = None
    for rel in rels_root:
        if rel.get("Id") == rid:
            target = rel.get("Target")
            break
    if target is None:
        raise ValueError(f"Relationship '{rid}' not found in workbook.xml.rels")

    target = target.lstrip("/")
    return target if target.startswith("xl/") else "xl/" + target


def add_column_via_zip_surgery(
    file_bytes: bytes,
    sheet_name: str,
    new_col_idx: int,
    new_col_name: str,
    row_values: List[str],
) -> Tuple[bytes, int]:
    """
    Appends one new column of string values to an existing sheet by
    editing only that sheet's raw XML inside the .xlsx zip.
    """
    col_letter = _col_letters(new_col_idx)
    header_cell = f'<c r="{col_letter}1" t="inlineStr"><is><t xml:space="preserve">{_escape_xml_text(new_col_name)}</t></is></c>'

    values_by_row = {i + 2: v for i, v in enumerate(row_values)}
    rows_written = set()

    def make_cell(row_num: int) -> str:
        val = values_by_row.get(row_num)
        if val is None:
            return ""
        rows_written.add(row_num)
        ref = f"{col_letter}{row_num}"
        return f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{_escape_xml_text(val)}</t></is></c>'

    def replace_row(m: re.Match) -> str:
        row_num = int(m.group(1))
        attrs = m.group(2)
        closer = m.group(3)
        if row_num == 1:
            new_cell = header_cell
        else:
            new_cell = make_cell(row_num)
            if not new_cell:
                return m.group(0)

        if closer == "/>":
            return f'<row r="{row_num}"{attrs}>{new_cell}</row>'
        else:
            inner = m.group(4) or ""
            return f'<row r="{row_num}"{attrs}>{inner}{new_cell}</row>'

    buf_in = io.BytesIO(file_bytes)
    with zipfile.ZipFile(buf_in, "r") as zin:
        sheet_path = _find_sheet_xml_path(zin, sheet_name)
        sheet_xml = zin.read(sheet_path).decode("utf-8")
        new_sheet_xml = _ROW_RE.sub(replace_row, sheet_xml)

        expected_rows = set(values_by_row.keys())
        if not expected_rows.issubset(rows_written):
            raise ValueError(
                f"{len(expected_rows - rows_written)} data row(s) were not found as "
                "<row> elements in the sheet XML, falling back to safe method."
            )

        buf_out = io.BytesIO()
        with zipfile.ZipFile(buf_out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = new_sheet_xml.encode("utf-8") if item.filename == sheet_path else zin.read(item.filename)
                zout.writestr(item, data)

    buf_out.seek(0)
    return buf_out.getvalue(), len(rows_written)


def get_sheet_names(file_bytes: bytes) -> List[str]:
    return pd.ExcelFile(io.BytesIO(file_bytes)).sheet_names


def get_columns(file_bytes: bytes, sheet_name: str) -> List[str]:
    return [str(c) for c in pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name, nrows=0).columns]


def get_row_count(file_bytes: bytes, sheet_name: str) -> int:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
    ws = wb[sheet_name]
    count = ws.max_row - 1 if ws.max_row else 0
    wb.close()
    return max(count, 0)


def process_lp_delete(
    file_bytes: bytes,
    sheet_edit: str,
    sheet_keywords: str,
    target_col: str,
    new_col_name: str,
    whole_word: bool = True,
    progress_callback: Optional[Callable[[float, str], None]] = None,
    log_callback: Optional[Callable[[str], None]] = None,
    is_cancelled: Optional[Callable[[], bool]] = None,
) -> Dict[str, Any]:
    """
    Execute full LP_Delete workflow and return results dictionary.
    """
    def log(msg: str):
        if log_callback:
            log_callback(msg)

    def progress(p: float, text: str = ""):
        if progress_callback:
            progress_callback(p, text)

    t0 = time.time()
    columns = get_columns(file_bytes, sheet_edit)
    row_count = get_row_count(file_bytes, sheet_edit)

    log(f"Loading column '{target_col}' from sheet '{sheet_edit}' ({row_count:,} rows)...")
    progress(0.05, "Loading target column...")

    search_col = pd.read_excel(
        io.BytesIO(file_bytes), sheet_name=sheet_edit, usecols=[target_col]
    )[target_col].fillna("").astype(str)

    log(f"Loading keywords from sheet '{sheet_keywords}'...")
    progress(0.12, "Loading keywords...")
    df_keywords_raw = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_keywords, header=None)
    keywords = df_keywords_raw.iloc[:, 0].dropna().astype(str).str.strip().tolist()
    keywords = [kw for kw in keywords if kw]
    log(f"Found {len(keywords):,} keyword(s).")

    if not keywords:
        log("No keywords found, all rows marked 'No'.")
        match_list = ["No"] * len(search_col)
        match_count = 0
    else:
        keywords_escaped = sorted((re.escape(kw) for kw in keywords), key=len, reverse=True)
        boundary = r"\b" if whole_word else ""
        pattern = boundary + "(?:" + "|".join(keywords_escaped) + ")" + boundary

        total_rows = len(search_col)
        chunk_size = 20000
        num_chunks = max(1, math.ceil(total_rows / chunk_size))
        results = []
        running_matches = 0
        scan_start = time.time()

        log(f"Scanning {total_rows:,} rows across {num_chunks} chunk(s)...")

        for i in range(num_chunks):
            if is_cancelled and is_cancelled():
                log("Process stopped by user.")
                return {"cancelled": True}

            start_idx = i * chunk_size
            end_idx = min((i + 1) * chunk_size, total_rows)
            chunk = search_col.iloc[start_idx:end_idx]
            match_mask = chunk.str.contains(pattern, case=False, regex=True, na=False)
            results.append(match_mask)
            running_matches += int(match_mask.sum())

            pct = 0.15 + 0.60 * ((i + 1) / num_chunks)
            elapsed = time.time() - scan_start
            rate = end_idx / elapsed if elapsed > 0 else 0
            progress(pct, f"Scanning chunk {i + 1}/{num_chunks}...")
            log(f"  Chunk {i + 1}/{num_chunks} (rows {start_idx:,}-{end_idx:,}): {running_matches:,} matches so far (~{rate:,.0f} rows/s)")

        final_mask = pd.concat(results)
        match_count = int(final_mask.sum())
        match_list = final_mask.map({True: "Yes", False: "No"}).tolist()
        log(f"Scanning complete in {time.time() - scan_start:.1f}s, {match_count:,} total matches.")

    new_col_idx = len(columns) + 1
    write_start = time.time()
    progress(0.78, "Writing updated column into workbook...")

    def write_via_openpyxl() -> bytes:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        ws = wb[sheet_edit]
        ws.cell(row=1, column=new_col_idx, value=new_col_name)
        for idx, value in enumerate(match_list):
            ws.cell(row=idx + 2, column=new_col_idx, value=value)
        out = io.BytesIO()
        wb.save(out)
        wb.close()
        out.seek(0)
        return out.getvalue()

    used_fast_path = False
    log("Attempting fast XML zip surgery...")
    try:
        output_bytes, rows_written = add_column_via_zip_surgery(
            file_bytes, sheet_edit, new_col_idx, new_col_name, match_list
        )
        used_fast_path = True
        log(f"Fast XML path completed: {rows_written:,} rows in {time.time() - write_start:.1f}s.")
    except Exception as e:
        log(f"Fast path bypassed ({e}), writing with standard workbook engine...")
        output_bytes = write_via_openpyxl()
        log(f"Saved via fallback in {time.time() - write_start:.1f}s.")

    # Verification
    log("Verifying written output matches computed results...")
    progress(0.92, "Verifying output...")
    verify_values = pd.read_excel(
        io.BytesIO(output_bytes), sheet_name=sheet_edit, usecols=[new_col_idx - 1]
    ).iloc[:, 0].tolist()

    verification_ok = (
        len(verify_values) == len(match_list)
        and all(a == b for a, b in zip(verify_values, match_list))
    )

    if not verification_ok and used_fast_path:
        log("Fast path verification discrepancy, re-saving with trusted openpyxl engine...")
        output_bytes = write_via_openpyxl()
        verify_values = pd.read_excel(
            io.BytesIO(output_bytes), sheet_name=sheet_edit, usecols=[new_col_idx - 1]
        ).iloc[:, 0].tolist()
        verification_ok = (
            len(verify_values) == len(match_list)
            and all(a == b for a, b in zip(verify_values, match_list))
        )

    if not verification_ok:
        raise ValueError("Verification failed: written workbook does not match expected computed results.")

    log(f"Verification passed ({len(verify_values):,} rows checked).")
    progress(1.0, "Completed!")

    # Generate preview dataframe (first 50 rows)
    preview_df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_edit, nrows=50)
    preview_df[new_col_name] = match_list[:50]
    preview_records = preview_df.fillna("").to_dict(orient="records")

    return {
        "cancelled": False,
        "match_count": match_count,
        "total_rows": len(search_col),
        "elapsed_seconds": round(time.time() - t0, 2),
        "output_bytes": output_bytes,
        "preview_records": preview_records,
        "preview_columns": list(preview_df.columns),
    }
