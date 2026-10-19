"""
core/lp_datewise.py, LP_DateWiseData stratified random sampling engine across 15-minute windows.
"""

import io
import math
import random
import time
from typing import Any, Callable, Dict, List, Optional, Tuple

import openpyxl
import pandas as pd


def get_sheet_names(file_bytes: bytes, filename: str) -> Optional[List[str]]:
    if filename.lower().endswith(".csv"):
        return None
    try:
        return pd.ExcelFile(io.BytesIO(file_bytes)).sheet_names
    except Exception:
        return None


def get_columns(file_bytes: bytes, filename: str, sheet_name: Optional[str] = None) -> List[str]:
    if filename.lower().endswith(".csv"):
        return [str(c) for c in pd.read_csv(io.BytesIO(file_bytes), nrows=0).columns]
    return [str(c) for c in pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name, nrows=0).columns]


def load_dataframe(file_bytes: bytes, filename: str, sheet_name: Optional[str] = None) -> pd.DataFrame:
    if filename.lower().endswith(".csv"):
        return pd.read_csv(io.BytesIO(file_bytes))
    target_sheet = sheet_name if sheet_name is not None else 0
    return pd.read_excel(io.BytesIO(file_bytes), sheet_name=target_sheet)


def process_lp_datewise(
    file_bytes: bytes,
    filename: str,
    datetime_col: str,
    target_total: int = 10000,
    result_column_name: str = "Selected",
    seed_input: str = "",
    sheet_name: Optional[str] = None,
    progress_callback: Optional[Callable[[float, str], None]] = None,
    log_callback: Optional[Callable[[str], None]] = None,
    is_cancelled: Optional[Callable[[], bool]] = None,
) -> Dict[str, Any]:
    """
    Execute stratified 15-minute sampling across active working hours (10:00-18:00).
    """
    def log(msg: str):
        if log_callback:
            log_callback(msg)

    def progress(p: float, text: str = ""):
        if progress_callback:
            progress_callback(p, text)

    t0 = time.time()
    log(f"Loading dataset from '{filename}'...")
    progress(0.05, "Reading file...")

    df_raw = load_dataframe(file_bytes, filename, sheet_name)
    total_raw_rows = len(df_raw)
    log(f"Loaded {total_raw_rows:,} rows and {len(df_raw.columns)} columns.")

    rng = random.Random(int(seed_input)) if seed_input.strip() else random.Random()
    df = df_raw.copy()

    log(f"Parsing date/time column '{datetime_col}'...")
    progress(0.10, "Parsing dates & times...")
    df["_dt"] = pd.to_datetime(df[datetime_col], errors="coerce")

    n_bad = int(df["_dt"].isna().sum())
    if n_bad:
        log(f"{n_bad:,} rows had unparseable dates/times and are excluded.")

    df_valid = df.dropna(subset=["_dt"]).copy()
    start_time = pd.Timestamp("10:00:00").time()
    end_time = pd.Timestamp("18:00:00").time()
    window_mask = df_valid["_dt"].dt.time.apply(lambda t: start_time <= t <= end_time)
    n_outside_window = int((~window_mask).sum())
    if n_outside_window:
        log(f"{n_outside_window:,} rows outside 10:00–18:00 window are excluded.")

    df_valid = df_valid.loc[window_mask].copy()
    if df_valid.empty:
        raise ValueError("No rows remained after filtering to the 10:00–18:00 working window.")

    df_valid["_date"] = df_valid["_dt"].dt.date
    df_valid["_hour"] = df_valid["_dt"].dt.hour
    df_valid["_minute"] = df_valid["_dt"].dt.minute
    df_valid["_quarter"] = df_valid["_minute"] // 15

    progress(0.18, "Structuring quarter-hour windows...")
    days = sorted(df_valid["_date"].unique())
    n_days = len(days)
    log(f"Found {n_days} distinct day(s): {days[0]} to {days[-1]}.")

    groups = df_valid.groupby(["_date", "_hour", "_quarter"]).indices
    log(f"Found {len(groups):,} active quarter-hour window buckets.")

    # Target calculation
    per_day_target = target_total / n_days
    quarter_targets: Dict[Tuple[Any, int, int], float] = {}
    hours_by_day: Dict[Any, set] = {}
    quarters_by_dh: Dict[Tuple[Any, int], List[int]] = {}
    for (d, h, q) in groups.keys():
        hours_by_day.setdefault(d, set()).add(h)
        quarters_by_dh.setdefault((d, h), []).append(q)
    for d, day_hours in hours_by_day.items():
        per_hour_target = per_day_target / len(day_hours)
        for h in day_hours:
            qs = quarters_by_dh[(d, h)]
            per_quarter_target = per_hour_target / len(qs)
            for q in qs:
                quarter_targets[(d, h, q)] = per_quarter_target

    quarter_keys = list(groups.keys())
    capacity = {k: len(groups[k]) for k in quarter_keys}

    # Floor + largest-remainder allocation
    raw_targets = {k: quarter_targets.get(k, 0.0) for k in quarter_keys}
    base_alloc = {k: min(int(math.floor(v)), capacity[k]) for k, v in raw_targets.items()}
    remainder_pool = target_total - sum(base_alloc.values())

    fractions = sorted(
        quarter_keys,
        key=lambda k: (raw_targets[k] - math.floor(raw_targets[k])),
        reverse=True,
    )
    i = 0
    guard = 0
    while remainder_pool > 0 and guard < 10 * len(quarter_keys) + 10:
        k = fractions[i % len(fractions)]
        if base_alloc[k] < capacity[k]:
            base_alloc[k] += 1
            remainder_pool -= 1
        i += 1
        guard += 1

    # Shortfall redistribution
    shortfall = target_total - sum(base_alloc.values())
    if shortfall > 0:
        log(f"Redistributing shortfall of {shortfall:,} records to windows with spare capacity...")
        spare_keys = [k for k in quarter_keys if base_alloc[k] < capacity[k]]
        rng.shuffle(spare_keys)
        idx = 0
        guard = 0
        while shortfall > 0 and spare_keys and guard < 20 * len(spare_keys) + 20:
            k = spare_keys[idx % len(spare_keys)]
            if base_alloc[k] < capacity[k]:
                base_alloc[k] += 1
                shortfall -= 1
            idx += 1
            guard += 1
            spare_keys = [k for k in spare_keys if base_alloc[k] < capacity[k]] or spare_keys

    # Actual random sampling
    progress(0.30, "Sampling within 15-minute buckets...")
    selected_index_set = set()
    total_windows = len(quarter_keys)

    for progress_i, k in enumerate(quarter_keys):
        if is_cancelled and is_cancelled():
            log("Sampling cancelled by user.")
            return {"cancelled": True}

        idx_positions = groups[k]
        take_n = base_alloc[k]
        if take_n > 0:
            chosen_positions = rng.sample(list(idx_positions), min(take_n, len(idx_positions)))
            chosen_labels = df_valid.index[chosen_positions]
            selected_index_set.update(chosen_labels)

        if progress_i % max(1, total_windows // 40) == 0 or progress_i == total_windows - 1:
            frac = 0.30 + 0.55 * (progress_i + 1) / total_windows
            progress(frac, f"Sampling window {progress_i + 1}/{total_windows}...")

    log(f"Total records selected: {len(selected_index_set):,} (target was {target_total:,}).")
    progress(0.88, "Building final output workbook...")

    df_out = df.drop(columns=["_dt"], errors="ignore")
    if result_column_name in df_out.columns:
        df_out = df_out.drop(columns=[result_column_name])
    df_out[result_column_name] = 0
    df_out.loc[list(selected_index_set), result_column_name] = 1

    # Daily summary table
    summary_df = (
        df_valid.assign(_selected=df_valid.index.isin(selected_index_set))
        .groupby("_date")["_selected"]
        .sum()
        .reset_index()
        .rename(columns={"_date": "Date", "_selected": "Records Selected"})
    )
    summary_df["Date"] = summary_df["Date"].astype(str)
    summary_records = summary_df.to_dict(orient="records")

    # Generate output file bytes preserving other sheets
    out_buffer = io.BytesIO()
    if filename.lower().endswith(".csv"):
        df_out.to_excel(out_buffer, index=False, engine="openpyxl")
    else:
        # ponytail: workbook parsed twice (pandas + openpyxl); optimise for large multi-sheet files
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        ws = wb[sheet_name] if sheet_name is not None else wb.worksheets[0]
        # Old merges/filters/tables would point at stale ranges after the rewrite
        for rng in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rng))
        ws.auto_filter.ref = None
        for tbl in list(ws.tables):
            del ws.tables[tbl]
        ws.delete_rows(1, ws.max_row)
        ws.append(list(df_out.columns))
        for row in df_out.astype(object).where(df_out.notna(), None).itertuples(index=False, name=None):
            ws.append(list(row))
        wb.calculation.fullCalcOnLoad = True  # openpyxl drops cached formula values; Excel recalcs on open
        wb.save(out_buffer)
        wb.close()

    out_buffer.seek(0)
    output_bytes = out_buffer.getvalue()

    preview_records = df_out.head(50).fillna("").to_dict(orient="records")
    progress(1.0, "Done!")

    return {
        "cancelled": False,
        "total_selected": len(selected_index_set),
        "total_rows": len(df_out),
        "target_total": target_total,
        "time_taken": round(time.time() - t0, 2),
        "summary_records": summary_records,
        "preview_columns": list(df_out.columns),
        "preview_records": preview_records,
        "output_bytes": output_bytes,
    }
