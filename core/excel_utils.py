"""
core/excel_utils.py, Comprehensive Excel styling and workbook management for Policy Reader.
Matches DailyTasks formatting, auto-fit column widths, openpyxl template mapping,
formula generation, date parsing, and frozen header styling (#D9E1F2).
"""

import logging
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent.parent
DEMO_COLUMNS_WORKBOOK = BASE_DIR / "DemoColumns.xlsx"
DEFAULT_POLICY_OUTPUT = BASE_DIR / "policy_reading_outputs.xlsx"

EXCEL_DATE_NUMBER_FORMAT = "DD-MM-YYYY"

DATE_TEMPLATE_COLUMNS: Set[str] = {
    "PolicyIssueDate",
    "PolicyStartDate",
    "PolicyEndDate",
    "TPStartDate",
    "TPEndDate",
    "DOB",
    "TranDated",
}

NUMERIC_FIELDS: Set[str] = {
    "IDV",
    "OD_NetPremium",
    "TP_NetPremium",
    "GST",
    "NetPremium",
    "GrossPremium",
    "CurrentNCB_%",
    "PreviousNCB_%",
    "Year_Of_Manufacturing",
    "CC/KWH",
    "Passenger",
    "GVW",
    "ChequeAmount",
    "PaymentAmount",
    "Pin",
}

FALLBACK_TEMPLATE_POLICY_COLUMNS: List[str] = [
    "Sno",
    "BatchId",
    "FileName",
    "Customer_Name",
    "Address",
    "City",
    "State",
    "Pin",
    "PhoneNo",
    "MobileNo",
    "DOB",
    "Email",
    "Insurer",
    "BusinessType",
    "PolicyType",
    "VehicleNo",
    "Make",
    "Model & Variant",
    "Year_Of_Manufacturing",
    "ChassisNo",
    "EngineNo",
    "CC/KWH",
    "Fuel",
    "RTOName",
    "CurrentNCB_%",
    "PreviousNCB_%",
    "Passenger",
    "GVW",
    "PolicyIssueDate",
    "PolicyStartDate",
    "PolicyEndDate",
    "TPStartDate",
    "TPEndDate",
    "PolicyNo",
    "CPA",
    "IDV",
    "OD_NetPremium",
    "TP_NetPremium",
    "NetPremium",
    "GST",
    "GrossPremium",
    "Ops Remarks",
    "PaymentMode",
    "ChequeNumber",
    "BankName",
    "PaymentAmount",
    "TranDated",
]

REMOVED_TEMPLATE_HEADERS: Set[str] = {
    "refercode",
    "financier",
    "financername",
    "financer_name",
}

REPLACED_TEMPLATE_HEADERS: Dict[str, str] = {
    "txndate": "TranDated",
}

TEMPLATE_COLUMN_MAP: Dict[str, str] = {
    "sno": "Sno",
    "batch_id": "BatchId",
    "file_name": "FileName",
    "customer_name": "Customer_Name",
    "address": "Address",
    "city": "City",
    "state": "State",
    "pin": "Pin",
    "phone_number": "PhoneNo",
    "mobile_number": "MobileNo",
    "dob": "DOB",
    "email": "Email",
    "insurer": "Insurer",
    "business_type": "BusinessType",
    "policy_type": "PolicyType",
    "vehicle_number": "VehicleNo",
    "make": "Make",
    "model_variant": "Model & Variant",
    "year_of_manufacture": "Year_Of_Manufacturing",
    "chassis_number": "ChassisNo",
    "engine_number": "EngineNo",
    "cc_kwh": "CC/KWH",
    "fuel_type": "Fuel",
    "rto_name": "RTOName",
    "current_ncb": "CurrentNCB_%",
    "previous_ncb": "PreviousNCB_%",
    "passenger": "Passenger",
    "gvw": "GVW",
    "policy_issue_date": "PolicyIssueDate",
    "policy_start": "PolicyStartDate",
    "policy_end": "PolicyEndDate",
    "tp_start": "TPStartDate",
    "tp_end": "TPEndDate",
    "txn_date": "TranDated",
    "policy_number": "PolicyNo",
    "cpa": "CPA",
    "idv": "IDV",
    "od_net_premium": "OD_NetPremium",
    "tp_net_premium": "TP_NetPremium",
    "gst": "GST",
    "gross_premium": "GrossPremium",
    "net_premium": "NetPremium",
    "ops_remarks": "Ops Remarks",
    "payment_mode": "PaymentMode",
    "cheque_number": "ChequeNumber",
    "bank_name": "BankName",
    "payment_amount": "PaymentAmount",
}


def _clean_text(value: Any) -> str:
    return str(value or "").strip()


def _strip_country_code(value: Any) -> str:
    text = _clean_text(value)
    if not text:
        return ""

    # Check if phone number is masked (contains X, x, *, •)
    if re.search(r"[xX*•]", text):
        # Extract digits, masking chars, and leading '+'
        cleaned = re.sub(r"[^\d+xX*•]", "", text)
        if cleaned.startswith("+91"):
            cleaned = cleaned[3:]
        elif cleaned.startswith("+"):
            cleaned = cleaned[1:]
        elif cleaned.startswith("0091"):
            cleaned = cleaned[4:]
        elif cleaned.startswith("00"):
            cleaned = cleaned[2:]

        # Remove any lingering non-phone characters
        cleaned = re.sub(r"[^\dxX*•]", "", cleaned)
        if not cleaned:
            return ""

        # Handle country code/prefix if present
        if len(cleaned) == 12 and cleaned.startswith("91"):
            return cleaned[2:]
        if len(cleaned) == 11 and cleaned.startswith("0"):
            return cleaned[1:]
        if len(cleaned) > 10 and (cleaned.startswith("91") or text.startswith("+")):
            return cleaned[-10:]
        return cleaned

    # Not masked - process as normal number
    cleaned = re.sub(r"[^\d+]", "", text)
    if cleaned.startswith("+91"):
        cleaned = cleaned[3:]
    elif cleaned.startswith("+"):
        cleaned = cleaned[1:]
    elif cleaned.startswith("0091"):
        cleaned = cleaned[4:]
    elif cleaned.startswith("00"):
        cleaned = cleaned[2:]

    digits = re.sub(r"\D+", "", cleaned)
    if not digits:
        return ""
    if len(digits) == 12 and digits.startswith("91"):
        return digits[2:]
    if len(digits) == 11 and digits.startswith("0"):
        return digits[1:]
    if len(digits) > 10 and (digits.startswith("91") or text.startswith("+")):
        return digits[-10:]
    return digits if digits else text


def _derive_make_and_model_variant(source: dict) -> Tuple[str, str]:
    make = _clean_text(source.get("make"))
    make_model = _clean_text(source.get("make_model") or source.get("Make & Model"))
    model = _clean_text(source.get("model"))
    variant = _clean_text(source.get("variant"))
    model_variant = _clean_text(source.get("model_variant") or source.get("Model & Variant"))

    if not make and make_model:
        parts = make_model.split()
        if parts:
            make = parts[0]
            if not model:
                model = " ".join(parts[1:]).strip()

    if not model_variant:
        if model and variant:
            model_variant = f"{model} {variant}".strip()
        elif model:
            model_variant = model
        elif variant:
            model_variant = variant
        elif make_model:
            model_variant = make_model

    return make, model_variant


def _norm_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _normalize_header_key(value: Any) -> str:
    return _norm_header(value)


def _ensure_address_split_headers(headers: List[str]) -> List[str]:
    updated = list(headers)
    normalized_existing = set()
    for header in updated:
        normalized = _norm_header(header)
        if normalized in REMOVED_TEMPLATE_HEADERS:
            continue
        normalized_existing.add(normalized)

    normalized_to_header = {_norm_header(h): h for h in FALLBACK_TEMPLATE_POLICY_COLUMNS}
    normalized_to_header.update({_norm_header(v): v for v in REPLACED_TEMPLATE_HEADERS.values()})

    normalized_target: List[str] = []
    seen = set()

    for header in FALLBACK_TEMPLATE_POLICY_COLUMNS:
        key = _norm_header(header)
        if key not in seen:
            normalized_target.append(key)
            seen.add(key)

    for header in updated:
        key = _norm_header(header)
        if not key or key in REMOVED_TEMPLATE_HEADERS:
            continue
        key = _norm_header(REPLACED_TEMPLATE_HEADERS.get(key, normalized_to_header.get(key, header)))
        if key not in seen:
            normalized_target.append(key)
            seen.add(key)

    return [normalized_to_header.get(k, k) for k in normalized_target]


def load_template_policy_columns() -> List[str]:
    if DEMO_COLUMNS_WORKBOOK.exists():
        try:
            wb = openpyxl.load_workbook(DEMO_COLUMNS_WORKBOOK, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            headers = [
                str(ws.cell(row=1, column=col).value).strip()
                for col in range(1, ws.max_column + 1)
                if str(ws.cell(row=1, column=col).value or "").strip()
            ]
            headers = [
                header for header in headers
                if _norm_header(header) not in REMOVED_TEMPLATE_HEADERS
            ]
            headers = _ensure_address_split_headers(headers)
            wb.close()
            if headers:
                return headers
        except Exception:
            pass
    return _ensure_address_split_headers(list(FALLBACK_TEMPLATE_POLICY_COLUMNS))


TEMPLATE_POLICY_COLUMNS = load_template_policy_columns()


def _normalized_template_header_keys() -> List[str]:
    return [_normalize_header_key(header) for header in TEMPLATE_POLICY_COLUMNS]


def _template_header_alias_map() -> Dict[str, Set[str]]:
    aliases = {header: {_normalize_header_key(header)} for header in TEMPLATE_POLICY_COLUMNS}
    for source_key, template_header in TEMPLATE_COLUMN_MAP.items():
        aliases.setdefault(template_header, set()).add(_normalize_header_key(source_key))
    return aliases


TEMPLATE_HEADER_ALIASES = _template_header_alias_map()


def coerce_excel_date_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.date() if isinstance(value, datetime) else value
    if value is None:
        return None

    raw = str(value).strip()
    if not raw:
        return None

    iso_match = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$", raw)
    if iso_match:
        try:
            year, month, day = map(int, iso_match.groups())
            return date(year, month, day)
        except ValueError:
            return raw

    dmy_match = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})$", raw)
    if dmy_match:
        try:
            day, month, year_str = dmy_match.groups()
            day_val = int(day)
            month_val = int(month)
            year_val = int(year_str)
            if len(year_str) == 2:
                year_val += 2000 if year_val < 50 else 1900
            return date(year_val, month_val, day_val)
        except ValueError:
            return raw

    date_time_match = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})\s+\d{1,2}:\d{1,2}(?::\d{1,2})?$", raw)
    if date_time_match:
        try:
            day, month, year_str = date_time_match.groups()
            day_val = int(day)
            month_val = int(month)
            year_val = int(year_str)
            if len(year_str) == 2:
                year_val += 2000 if year_val < 50 else 1900
            return date(year_val, month_val, day_val)
        except ValueError:
            return raw

    return raw


def _coerce_excel_date_value(value: Any) -> Any:
    return coerce_excel_date_value(value)


def _map_row_to_template(row_data: dict) -> dict:
    mapped = {}
    normalized_source = {_normalize_header_key(k): v for k, v in row_data.items()}

    for template_header in TEMPLATE_POLICY_COLUMNS:
        aliases = TEMPLATE_HEADER_ALIASES.get(template_header, {_normalize_header_key(template_header)})
        val = ""
        for alias in aliases:
            if alias in normalized_source:
                val = normalized_source[alias]
                if val is not None and str(val).strip():
                    break
        mapped[template_header] = val if val is not None else ""

    if mapped.get("PhoneNo") or "phone_number" in row_data or "PhoneNo" in row_data:
        mapped["PhoneNo"] = _strip_country_code(mapped.get("PhoneNo") or row_data.get("phone_number") or row_data.get("PhoneNo"))
    if mapped.get("MobileNo") or "mobile_number" in row_data or "MobileNo" in row_data:
        mapped["MobileNo"] = _strip_country_code(mapped.get("MobileNo") or row_data.get("mobile_number") or row_data.get("MobileNo"))

    make_val, model_variant_val = _derive_make_and_model_variant(row_data)
    if make_val and not _clean_text(mapped.get("Make")):
        mapped["Make"] = make_val
    existing_mv = _clean_text(mapped.get("Model & Variant"))
    if model_variant_val and not existing_mv:
        mapped["Model & Variant"] = model_variant_val
    elif model_variant_val and existing_mv:
        if existing_mv.lower() in model_variant_val.lower():
            mapped["Model & Variant"] = model_variant_val
        elif model_variant_val.lower() not in existing_mv.lower():
            mapped["Model & Variant"] = f"{model_variant_val} {existing_mv}".strip()

    raw_tran = _clean_text(mapped.get("TranDated")) or _clean_text(row_data.get("txn_date") or row_data.get("TxnDate") or row_data.get("TranDated"))
    if "/" in raw_tran and re.search(r"[A-Za-z]", raw_tran):
        d_clean = re.search(r"\b(\d{2})(\d{2})(\d{4})\b", raw_tran)
        if d_clean:
            raw_tran = f"{d_clean.group(1)}/{d_clean.group(2)}/{d_clean.group(3)}"
        else:
            d_clean2 = re.search(r"(\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})", raw_tran)
            if d_clean2:
                raw_tran = d_clean2.group(1)
    mapped["TranDated"] = raw_tran

    if not mapped.get("FileName"):
        mapped["FileName"] = row_data.get("File Name", "") or row_data.get("file_name", "") or row_data.get("filename", "")

    mapped["NetPremium"] = 0
    mapped["GrossPremium"] = 0

    payment_mode_mapped = str(mapped.get("PaymentMode") or "").strip().lower()
    if payment_mode_mapped != "cheque":
        mapped["ChequeNumber"] = ""
        mapped["BankName"] = ""

    return mapped


def _normalize_template_row_values(template_row: dict) -> dict:
    normalized = dict(template_row)
    for header in DATE_TEMPLATE_COLUMNS:
        normalized[header] = coerce_excel_date_value(normalized.get(header, ""))
    return normalized


def _align_sheet_to_template_headers(ws) -> List[str]:
    existing_headers = [
        str(ws.cell(row=1, column=col).value) if ws.cell(row=1, column=col).value is not None else ""
        for col in range(1, ws.max_column + 1)
    ]
    if existing_headers == TEMPLATE_POLICY_COLUMNS:
        return existing_headers

    normalized_existing = {
        _normalize_header_key(header): idx
        for idx, header in enumerate(existing_headers, start=1)
        if str(header).strip()
    }

    remaining_headers = []
    for col_idx, header in enumerate(existing_headers, start=1):
        norm = _normalize_header_key(header)
        if norm in REMOVED_TEMPLATE_HEADERS:
            continue
        if norm in REPLACED_TEMPLATE_HEADERS:
            ws.cell(row=1, column=col_idx, value=REPLACED_TEMPLATE_HEADERS[norm])
            remaining_headers.append(REPLACED_TEMPLATE_HEADERS[norm])
        else:
            remaining_headers.append(header)

    for header in TEMPLATE_POLICY_COLUMNS:
        norm = _normalize_header_key(header)
        if norm not in normalized_existing and norm not in REMOVED_TEMPLATE_HEADERS:
            next_col = ws.max_column + 1
            ws.cell(row=1, column=next_col, value=header)
            remaining_headers.append(header)

    return [
        str(ws.cell(row=1, column=col).value) if ws.cell(row=1, column=col).value is not None else ""
        for col in range(1, ws.max_column + 1)
    ]


def _find_first_empty_data_row(ws, existing_headers: list) -> int:
    header_count = len(existing_headers)
    file_name_col = None
    for idx, header in enumerate(existing_headers, start=1):
        if _normalize_header_key(header) in {
            _normalize_header_key("FileName"),
            _normalize_header_key("File Name"),
            _normalize_header_key("file_name"),
            _normalize_header_key("filename"),
        }:
            file_name_col = idx
            break

    for row_idx in range(2, ws.max_row + 2):
        row_has_data = False
        for col_idx in range(1, header_count + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value not in (None, ""):
                row_has_data = True
                break

        if not row_has_data:
            return row_idx

        # A row with only the file name and no mapped policy fields is also a
        # reusable slot that should be filled from the top instead of writing to
        # the last appended row.
        if header_count >= 1 and file_name_col is not None:
            file_name_value = ws.cell(row=row_idx, column=file_name_col).value
            other_values_present = any(
                ws.cell(row=row_idx, column=col).value not in (None, "")
                for col in range(1, header_count + 1)
                if col != file_name_col
            )
            if file_name_value not in (None, "") and not other_values_present:
                return row_idx

    return ws.max_row + 1


def _trim_trailing_blank_rows(ws) -> None:
    while ws.max_row > 1:
        row_idx = ws.max_row
        row_is_blank = all(
            ws.cell(row=row_idx, column=col).value in (None, "")
            for col in range(1, ws.max_column + 1)
        )
        if row_is_blank:
            ws.delete_rows(row_idx, 1)
        else:
            break


def _ensure_policy_template_copy(output_excel: str | os.PathLike | None = None) -> Path:
    output_path = Path(output_excel or DEFAULT_POLICY_OUTPUT)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if not output_path.exists():
        wb = Workbook()
        ws = wb.active
        if ws is None:
            raise RuntimeError("Failed to create active worksheet in new workbook")
        ws.title = "Policies"
        ws.append(TEMPLATE_POLICY_COLUMNS)
        wb.save(output_path)
        wb.close()
        return output_path

    wb = load_workbook(output_path)
    if "Policies" in wb.sheetnames:
        ws = wb["Policies"]
    else:
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet("Policies")
        else:
            ws.title = "Policies"

    if ws is None:
        raise RuntimeError("Failed to access or create 'Policies' worksheet")

    header_has_values = any(
        str(ws.cell(row=1, column=col).value or "").strip()
        for col in range(1, ws.max_column + 1)
    )
    if not header_has_values:
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
        ws.append(TEMPLATE_POLICY_COLUMNS)
    else:
        _align_sheet_to_template_headers(ws)
        _trim_trailing_blank_rows(ws)

    wb.save(output_path)
    wb.close()
    return output_path


def _append_row_to_excel(output_excel: str | os.PathLike, sheet_name: str, row_data: dict) -> None:
    output_path = _ensure_policy_template_copy(output_excel)
    wb = load_workbook(output_path)

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(title=sheet_name)

    if ws is None:
        wb.close()
        raise RuntimeError(f"Failed to access or create worksheet '{sheet_name}'")

    template_row = _normalize_template_row_values(_map_row_to_template(row_data))
    if not ws._cells:
        existing_headers = list(TEMPLATE_POLICY_COLUMNS)
        ws.append(existing_headers)
    else:
        existing_headers = [
            str(ws.cell(row=1, column=col).value) if ws.cell(row=1, column=col).value is not None else ""
            for col in range(1, ws.max_column + 1)
        ]

    if ws._cells and not any(str(ws.cell(row=1, column=col).value).strip() for col in range(1, ws.max_column + 1)):
        sheet_position = wb.index(ws)
        wb.remove(ws)
        ws = wb.create_sheet(title=sheet_name, index=sheet_position)
        existing_headers = list(TEMPLATE_POLICY_COLUMNS)
        ws.append(existing_headers)
    elif ws._cells and existing_headers != TEMPLATE_POLICY_COLUMNS:
        existing_headers = _align_sheet_to_template_headers(ws)

    _trim_trailing_blank_rows(ws)

    target_row = _find_first_empty_data_row(ws, existing_headers)
    curr_fn = str(template_row.get("FileName") or "").strip()

    if not str(template_row.get("Sno") or "").strip():
        template_row["Sno"] = target_row - 1

    od_col_idx = None
    tp_col_idx = None
    gst_col_idx = None
    net_col_idx = None
    gross_col_idx = None

    for idx, header in enumerate(existing_headers, start=1):
        header_normalized = _normalize_header_key(header)
        if header_normalized == _normalize_header_key("OD_NetPremium"):
            od_col_idx = idx
        elif header_normalized == _normalize_header_key("TP_NetPremium"):
            tp_col_idx = idx
        elif header_normalized == _normalize_header_key("GST"):
            gst_col_idx = idx
        elif header_normalized == _normalize_header_key("NetPremium"):
            net_col_idx = idx
        elif header_normalized == _normalize_header_key("GrossPremium"):
            gross_col_idx = idx

    values = [template_row.get(header, "") for header in existing_headers]
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=target_row, column=col_idx)
        if col_idx == net_col_idx and od_col_idx and tp_col_idx:
            od_letter = get_column_letter(od_col_idx)
            tp_letter = get_column_letter(tp_col_idx)
            cell.value = f"={od_letter}{target_row}+{tp_letter}{target_row}"
        elif col_idx == gross_col_idx and net_col_idx and gst_col_idx:
            net_letter = get_column_letter(net_col_idx)
            gst_letter = get_column_letter(gst_col_idx)
            cell.value = f"={net_letter}{target_row}+{gst_letter}{target_row}"
        else:
            if existing_headers[col_idx - 1] in NUMERIC_FIELDS:
                try:
                    if value in (None, "", "0", 0):
                        cell.value = 0
                    else:
                        clean_num = str(value).replace(",", "").strip()
                        cell.value = float(clean_num) if "." in clean_num else int(clean_num)
                except (ValueError, TypeError):
                    cell.value = 0
            else:
                # Sanitize text against Excel formula injection
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
                    cell.value = f"'{value}"
                else:
                    cell.value = value

        if existing_headers[col_idx - 1] in DATE_TEMPLATE_COLUMNS:
            cell.number_format = EXCEL_DATE_NUMBER_FORMAT

    try:
        if ws is not None:
            wb.active = ws
        wb.save(output_path)
    finally:
        wb.close()


def format_excel_file(excel_file: str | Path | None = None, sheet_name: Optional[str] = None) -> None:
    excel_path = Path(excel_file or DEFAULT_POLICY_OUTPUT)
    if not excel_path.exists():
        raise FileNotFoundError(f"Workbook not found: {excel_path}")

    wb = openpyxl.load_workbook(excel_path)
    target_sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(bold=True)

    if target_sheet.max_row > 0:
        for cell in target_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        target_sheet.freeze_panes = "A2"

        long_text_cols = {"coverage_summary", "exclusions", "review_reason", "ops_remarks", "ops remarks"}
        for col_idx, col_cells in enumerate(target_sheet.columns, 1):
            col_letter = get_column_letter(col_idx)
            header_val = str(target_sheet[f"{col_letter}1"].value or "")

            alignment = Alignment(vertical="top", wrap_text=False)
            width = 20

            if header_val in ("File Name", "FileName"):
                width = 35
            elif _norm_header(header_val) in {_norm_header(k) for k in long_text_cols}:
                width = 50
                alignment = Alignment(vertical="top", wrap_text=True)
            else:
                max_len = 0
                for cell in col_cells:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                width = min(max_len + 2, 30)

            target_sheet.column_dimensions[col_letter].width = max(width, 12)
            for cell in col_cells[1:]:
                if header_val in DATE_TEMPLATE_COLUMNS:
                    coerced = coerce_excel_date_value(cell.value)
                    if coerced != cell.value:
                        cell.value = coerced
                    if cell.value not in (None, ""):
                        cell.number_format = EXCEL_DATE_NUMBER_FORMAT
                elif header_val in NUMERIC_FIELDS and cell.value not in (None, ""):
                    val_str = str(cell.value).strip().replace(",", "")
                    if val_str:
                        try:
                            if "." in val_str:
                                cell.value = float(val_str)
                            else:
                                cell.value = int(val_str)
                        except ValueError:
                            pass
                cell.alignment = alignment

    if target_sheet is not None:
        wb.active = target_sheet
    wb.save(excel_path)
    wb.close()
